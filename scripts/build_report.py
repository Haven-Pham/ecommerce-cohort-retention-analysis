"""Build the Excel report and the executive slide deck from the notebook's exports.

Run after the notebook (it reads reports/tables/*.csv and reports/figures/*.png):

    python scripts/build_report.py
"""
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TABLES = ROOT / "reports" / "tables"
FIGURES = ROOT / "reports" / "figures"
XLSX = ROOT / "reports" / "Cohort_Retention_Profitability_Report.xlsx"
PPTX = ROOT / "reports" / "Cohort_Retention_Executive_Deck.pptx"

NAVY = RGBColor(0x0A, 0x1C, 0x33)
BLUE = RGBColor(0x1D, 0x6A, 0xE5)
SKY = RGBColor(0x38, 0xBD, 0xF8)
GREY = RGBColor(0x52, 0x6E, 0x8C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SHEETS = {
    "Yearly KPIs": "yearly_kpis",
    "Margin by discount": "margin_by_discount_band",
    "Margin by category": "margin_by_category",
    "Repeat vs one-time": "repeat_vs_one_time",
    "Segments": "segment_summary",
    "Income brackets": "income_profitability",
    "Discount bands": "discount_profitability",
    "Monthly retention": "monthly_retention",
    "Quarterly retention": "quarterly_retention",
    "Data quality": "data_quality_checks",
}
PERCENT_HINTS = ("margin", "share", "rate", "discount", "retention")


def build_excel() -> None:
    header_fill = PatternFill("solid", fgColor="0A1C33")
    with pd.ExcelWriter(XLSX, engine="openpyxl") as writer:
        for sheet, name in SHEETS.items():
            table = pd.read_csv(TABLES / f"{name}.csv", index_col=0)
            table.to_excel(writer, sheet_name=sheet)
            ws = writer.sheets[sheet]
            percent_sheet = name.endswith("retention") or name.startswith("margin_by")
            for col_idx, column in enumerate([table.index.name or ""] + list(table.columns), start=1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", wrap_text=True)
                ws.column_dimensions[get_column_letter(col_idx)].width = 16 if col_idx > 1 else 22
                if col_idx == 1:
                    continue
                is_percent = percent_sheet or any(h in str(column).lower() for h in PERCENT_HINTS)
                for row in range(2, len(table) + 2):
                    value = ws.cell(row=row, column=col_idx)
                    if isinstance(value.value, (int, float)):
                        value.number_format = "0.0%" if is_percent else "#,##0.00" if isinstance(value.value, float) else "#,##0"
            ws.freeze_panes = "B2"
    print(f"Wrote {XLSX.relative_to(ROOT)}")


def _text(slide, text, left, top, width, height, size=16, bold=False, color=NAVY):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.text = line
        para.font.size = Pt(size)
        para.font.bold = bold
        para.font.color.rgb = color
        para.space_after = Pt(6)
    return box


def _header(slide, title, kicker):
    bar = slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()
    _text(slide, kicker.upper(), 0.6, 0.35, 12, 0.4, size=12, bold=True, color=BLUE)
    _text(slide, title, 0.6, 0.7, 12, 0.9, size=26, bold=True)


def _figure_slide(prs, kicker, title, figure, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, title, kicker)
    slide.shapes.add_picture(str(FIGURES / figure), Inches(0.6), Inches(1.75), width=Inches(7.9))
    _text(slide, "\n".join(f"• {b}" for b in bullets), 8.8, 1.9, 4.1, 5, size=15, color=GREY)


def build_deck() -> None:
    kpis = pd.read_csv(TABLES / "yearly_kpis.csv", index_col=0)
    repeat = pd.read_csv(TABLES / "repeat_vs_one_time.csv", index_col=0)
    discounts = pd.read_csv(TABLES / "discount_profitability.csv", index_col=0)
    k22, k23 = kpis.loc[2022], kpis.loc[2023]
    rep = repeat.loc["Repeat (2+ orders)"]
    deep_loss = discounts.loc[["21–40%", ">40%"], "profit"].sum()

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    _text(slide, "CUSTOMER COHORT, RETENTION & PROFITABILITY REVIEW", 0.8, 2.0, 11.5, 0.5, size=14, bold=True, color=SKY)
    _text(slide, "Growth is real. Profit is not keeping up.", 0.8, 2.6, 11.5, 1.2, size=40, bold=True, color=WHITE)
    _text(slide, f"E-commerce retailer · 2020–2023 · {int(repeat['customers'].sum()):,} customers · "
                 f"{int(kpis['orders'].sum()):,} orders\n"
                 "Pham Hong Hieu (Haven) · github.com/Haven-Pham", 0.8, 4.2, 11.5, 1.2, size=16, color=WHITE)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "Executive summary", "Five things management should know")
    _text(slide, "\n".join([
        f"1. 2023 sales +{k23.sales / k22.sales - 1:.0%}, but profit {k23.profit / k22.profit - 1:.0%}: margin {k22.margin:.1%} → {k23.margin:.1%} in every category.",
        f"2. Repeat customers are {rep.customer_share:.0%} of customers but {rep.profit_share:.0%} of profit (${rep.avg_profit_per_customer:,.0f} each vs ${repeat.loc['One-time', 'avg_profit_per_customer']:,.0f}).",
        "3. Only ~2% of a cohort buys in any later month; the median gap to a 2nd purchase is 364 days.",
        f"4. Discounts above 20% lost ${-deep_loss / 1e3:,.0f}K over four years.",
        "5. Income does not separate profitable customers; behaviour does.",
    ]), 0.6, 1.8, 12, 5, size=20)

    _figure_slide(prs, "Profitability", "2023: sales up, margin down", "01_yearly_sales_profit_margin.png", [
        f"Sales ${k22.sales / 1e6:.2f}M → ${k23.sales / 1e6:.2f}M",
        f"Profit ${k22.profit / 1e3:,.0f}K → ${k23.profit / 1e3:,.0f}K",
        "Margin gap worth ≈ $115K at 2023 volume",
    ])
    _figure_slide(prs, "Root cause", "Not discounts, not mix: unit economics", "01b_margin_drop_diagnosis.png", [
        "Undiscounted margin 33.4% → 28.1%",
        "Every category lost 4–8 pp",
        "Category mix moved < 1 pp",
        "Next step: cost and list-price data",
    ])
    _figure_slide(prs, "Retention", "Most customers never come back", "03_monthly_retention_heatmap_2022.png", [
        "0.4–2.9% of a 2022 cohort active in any later month",
        "No decay curve: return is rare but steady",
        "Target the second purchase",
    ])
    _figure_slide(prs, "Customer value", "Repeat buyers carry the profit", "05_repeat_customer_value.png", [
        f"{rep.customer_share:.0%} of customers → {rep.profit_share:.0%} of profit",
        "Median 364 days to 2nd purchase",
        "+5% repeat conversion ≈ +$43K profit",
    ])
    _figure_slide(prs, "Pricing", "Deep discounts lose money", "08_discount_profit_leakage.png", [
        "No discount: 31.7% margin",
        "21–40%: −8.1% · >40%: −27.0%",
        f"${-deep_loss / 1e3:,.0f}K lost on deep-discount lines",
    ])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "Recommendations", "What to do next")
    _text(slide, "\n".join([
        "1. Finance · Procurement: explain the 2023 margin drop with cost and list-price data (≈ $115K/yr).",
        "2. Commercial: cap discounts at 20% with approval above it (up to ≈ $72K/yr).",
        "3. CRM: second-purchase programme timed at months 10–12, plus a tested 60–90-day nudge (≈ +$43K).",
        "4. Sales: account management for Corporate (30% of customers, 57% of profit).",
        "5. Marketing: replace income targeting with RFM and discount-sensitivity segments.",
    ]), 0.6, 1.8, 12, 5, size=20)

    prs.save(PPTX)
    print(f"Wrote {PPTX.relative_to(ROOT)} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build_excel()
    build_deck()
